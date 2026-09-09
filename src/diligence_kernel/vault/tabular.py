"""Spreadsheets and delimited files, read as tables rather than as prose.

`00a` classifies these as **Records** — cap tables, stock ledgers, employee censuses, loss
runs, schedules of subsidiaries — and notes that an as-of date is what makes a record
usable. They are input documents, not just the place arithmetic happens.

Flattening a sheet into prose and chunking it by paragraph loses two things that a reader
needs and cannot recover:

- **Column headers.** A chunk holding rows 27 to 52 of a census, with the header left behind
  in an earlier chunk, is a grid of unlabelled numbers. Nothing downstream can tell salary
  from bonus, and a model asked to read it will guess.
- **Sheet boundaries.** A cap table followed by a census, with no marker between them, reads
  as one table with a sudden change of shape.

So a table is chunked by rows, and every chunk restates the sheet name and the header. The
repetition is written into the stored text as well, the way a printed schedule repeats its
headings on each page — which keeps each chunk an exact slice of the document, so character
offsets and evidence provenance stay true.
"""

from __future__ import annotations

import csv
from dataclasses import dataclass, field
from pathlib import Path

TABULAR_SUFFIXES = frozenset({".xlsx", ".xlsm", ".csv", ".tsv"})
#: Rows to inspect when looking for the header row.
HEADER_SCAN_ROWS = 12
#: Target size of one chunk of rows, in characters.
CHUNK_CHARS = 2400


@dataclass(slots=True)
class Sheet:
    name: str
    preamble: list[str] = field(default_factory=list)
    header: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)

    @property
    def row_count(self) -> int:
        return len(self.rows)


def is_tabular(suffix: str) -> bool:
    return suffix.lower() in TABULAR_SUFFIXES


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _split_header(grid: list[list[str]]) -> tuple[list[str], list[str], list[list[str]]]:
    """Find the header row. Everything above it is preamble, everything below is data.

    A spreadsheet usually opens with a title and an as-of date before the grid starts. The
    header is taken to be the widest row in the opening band, which is what a title row and
    a stray `As of` row are not.
    """
    if not grid:
        return [], [], []
    widths = [sum(1 for c in row if c) for row in grid[:HEADER_SCAN_ROWS]]
    if not widths:
        return [], [], grid
    best = max(widths)
    if best < 2:
        return [], [], grid
    index = widths.index(best)
    preamble = ["\t".join(c for c in row if c) for row in grid[:index] if any(c for c in row)]
    return preamble, grid[index], grid[index + 1 :]


def read_workbook(path: Path) -> list[Sheet]:
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    sheets: list[Sheet] = []
    try:
        for worksheet in workbook.worksheets:
            grid = [[_cell(c) for c in row] for row in worksheet.iter_rows(values_only=True)]
            grid = [row for row in grid if any(row)]
            preamble, header, rows = _split_header(grid)
            sheets.append(Sheet(worksheet.title, preamble, header, rows))
    finally:
        workbook.close()
    return sheets


def read_delimited(path: Path) -> list[Sheet]:
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    with path.open(newline="", encoding="utf-8-sig", errors="replace") as handle:
        grid = [[_cell(c) for c in row] for row in csv.reader(handle, delimiter=delimiter)]
    grid = [row for row in grid if any(row)]
    preamble, header, rows = _split_header(grid)
    return [Sheet(path.stem, preamble, header, rows)]


def read_sheets(path: Path) -> list[Sheet]:
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        return read_workbook(path)
    return read_delimited(path)


def render(sheets: list[Sheet], *, chunk_chars: int = CHUNK_CHARS) -> tuple[str, list[str]]:
    """Render sheets to text, and to chunks that are exact slices of that text.

    Each chunk opens with the sheet name and the header row, so it can be read alone.
    """
    chunks: list[str] = []
    for sheet in sheets:
        header_line = "\t".join(sheet.header) if sheet.header else ""
        opening = [f"# Sheet: {sheet.name}"]
        if sheet.preamble:
            opening.extend(sheet.preamble)
        if header_line:
            opening.append(header_line)

        if not sheet.rows:
            # A sheet with no data contributes nothing to read. Emitting its name alone
            # would make a blank workbook look like content.
            if sheet.header or sheet.preamble:
                chunks.append("\n".join(opening))
            continue

        current: list[str] = list(opening)
        size = sum(len(line) + 1 for line in current)
        first = True
        for row in sheet.rows:
            line = "\t".join(row)
            if size + len(line) + 1 > chunk_chars and not first:
                chunks.append("\n".join(current))
                # A continuation repeats what a reader needs: which sheet, which columns.
                current = [f"# Sheet: {sheet.name} (continued)"]
                if header_line:
                    current.append(header_line)
                size = sum(len(x) + 1 for x in current)
            current.append(line)
            size += len(line) + 1
            first = False
        if len(current) > (2 if header_line else 1) or not chunks:
            chunks.append("\n".join(current))

    return "\n\n".join(chunks), chunks


def spans_for(full_text: str, chunks: list[str]) -> list[tuple[int, int]]:
    """Character spans of chunks joined by a blank line — exact by construction."""
    spans: list[tuple[int, int]] = []
    cursor = 0
    for piece in chunks:
        spans.append((cursor, cursor + len(piece)))
        cursor += len(piece) + 2
    return spans
